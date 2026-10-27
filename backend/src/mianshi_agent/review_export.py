"""Export the same structured report shown on the page as editable Word content."""

from io import BytesIO

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def export_review(record, display_name):
    if record["status"] != "completed" or not record["summary"]:
        raise RuntimeError("external_review_export_not_ready")
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2), Cm(1.9)
    section.left_margin = section.right_margin = Cm(2.2)
    for name, font, size, bold, before, after in [
        ("Normal", "宋体", 10.5, False, 0, 5),
        ("Title", "微软雅黑", 24, True, 11, 7),
        ("Subtitle", "微软雅黑", 11, False, 11, 7),
        ("Heading 1", "微软雅黑", 16, True, 15, 7),
        ("Heading 2", "微软雅黑", 13, True, 11, 7),
    ]:
        style = doc.styles[name]
        style.font.name = font
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.22 if name == "Normal" else 1.0
        if name.startswith("Heading"):
            style.paragraph_format.keep_with_next = True
    doc.core_properties.author = "Interview Agent"
    doc.core_properties.title = record["title"] + "面试复盘"
    footer = section.footer.paragraphs[0]
    footer.alignment = 1
    footer.add_run("第 ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页")
    doc.add_paragraph("本轮面试逐题复盘", "Title")
    doc.add_paragraph(f"{display_name}  {record['title']}", "Subtitle")
    doc.add_heading("本轮结论", 1)
    doc.add_paragraph(record["summary"]["conclusion"])
    doc.add_paragraph(
        f"本次完整处理 {record['coverage']['question_count']} 个问题及追问。"
        "实际回答整理与建议回答分别呈现，原文独立保留。"
    )
    sources = {}
    for index, question in enumerate(record["questions"], 1):
        result = question["result"]
        content = result["content"]
        doc.add_heading(f"{index} {question['session']} {question['topic'] or '面试问答'}", 1)
        doc.add_heading("面试问题", 2)
        doc.add_paragraph(question["question"])
        doc.add_heading("实际回答整理", 2)
        doc.add_paragraph(content["answer_clean"] or "该题回答未记录。")
        doc.add_paragraph(f"原文位置：字符 {question['start'] + 1} 至 {question['end']}。")
        doc.add_heading("逐题分析", 2)
        doc.add_paragraph(content["analysis"])
        for quote in content["quotes"]:
            doc.add_paragraph("回答原文依据：“" + quote + "”")
        doc.add_heading("建议回答", 2)
        doc.add_paragraph(content["suggestion"])
        if content["explanation"]:
            doc.add_heading("补充说明", 2)
            doc.add_paragraph(content["explanation"])
        if content["gaps"]:
            doc.add_heading("待补信息", 2)
            for gap in content["gaps"]:
                doc.add_paragraph(gap, "List Bullet")
        if result["evidence"]:
            doc.add_heading("引用依据", 2)
            for source in result["evidence"]:
                sources[source["chunk_id"]] = source
                doc.add_paragraph(f"{source['source_title']}  {source['locator']}")
                doc.add_paragraph(source["text"])
    doc.add_heading("下一步准备重点", 1)
    for priority in record["summary"]["priorities"]:
        doc.add_paragraph(priority, "List Number")
    doc.add_heading("记录与证据边界", 1)
    for limitation in record["summary"]["limitations"]:
        doc.add_paragraph(limitation, "List Bullet")
    doc.add_heading("本次复盘使用的项目资料", 1)
    for source in sources.values():
        doc.add_paragraph(
            f"{source['source_title']}  {source['locator']}  ID {source['chunk_id']}", "List Bullet"
        )
    if not sources:
        doc.add_paragraph("本次未引用项目资料；个人贡献、实际效果仍需独立依据。")
    output = BytesIO()
    doc.save(output)
    return output.getvalue()
